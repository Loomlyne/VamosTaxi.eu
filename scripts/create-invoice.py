from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "deliverables"
OUT.mkdir(exist_ok=True)
OUTPUT = OUT / "invoice-vamos-taxi-eur-1500.docx"

YELLOW = "FDC20B"
CHARCOAL = "1E1F1F"
LIGHT_GRAY = "F4F4F2"
MUTED = "666666"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, values in kwargs.items():
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in values.items():
            element.set(qn("w:{}".format(key)), str(value))


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid_cols = table._tbl.tblGrid.gridCol_lst
    for col, width in zip(grid_cols, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def font(run, size=10.5, color=CHARCOAL, bold=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def para(cell_or_doc, text="", size=10.5, color=CHARCOAL, bold=False, align=None, after=0, before=0):
    p = cell_or_doc.add_paragraph() if hasattr(cell_or_doc, "add_paragraph") else cell_or_doc.paragraphs[0]
    if p.runs:
        p.clear()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    font(r, size, color, bold)
    return p


def cell_text(cell, text, size=10.5, color=CHARCOAL, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.clear()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = align
    r = p.add_run(text)
    font(r, size, color, bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    return p


def add_label_value(doc, label, value, y):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(4.6)
    p.paragraph_format.first_line_indent = Inches(-1.0)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(f"{label}  ")
    font(r, 9, MUTED, True)
    r = p.add_run(value)
    font(r, 10, CHARCOAL, False)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.right_margin = Inches(0.72)
section.bottom_margin = Inches(0.6)
section.left_margin = Inches(0.72)
section.header_distance = Inches(0.25)
section.footer_distance = Inches(0.25)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)

# Header
header = section.header
header_p = header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header_p.paragraph_format.space_after = Pt(0)
run = header_p.add_run("VAMOS TAXI  |  PROJECT INVOICE")
font(run, 8.5, MUTED, True)

# Top title row
top = doc.add_table(rows=1, cols=2)
top.alignment = WD_TABLE_ALIGNMENT.LEFT
top.autofit = False
set_table_widths(top, [5400, 3960])
for cell in top.rows[0].cells:
    set_cell_border(cell, top={"val": "nil"}, bottom={"val": "nil"}, start={"val": "nil"}, end={"val": "nil"})

left, right = top.rows[0].cells
cell_text(left, "Koussay Zayani", 18, CHARCOAL, True)
p = left.add_paragraph()
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Independent product design and development")
font(r, 10, MUTED)

cell_text(right, "INVOICE", 26, CHARCOAL, True, WD_ALIGN_PARAGRAPH.RIGHT)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

# Stripe
stripe = doc.add_table(rows=1, cols=1)
stripe.alignment = WD_TABLE_ALIGNMENT.LEFT
stripe.autofit = False
set_table_widths(stripe, [9360])
cell = stripe.cell(0, 0)
set_cell_shading(cell, YELLOW)
set_cell_margins(cell, top=55, bottom=55, start=0, end=0)
cell_text(cell, "", 1)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Billing information
info = doc.add_table(rows=1, cols=2)
info.alignment = WD_TABLE_ALIGNMENT.LEFT
info.autofit = False
set_table_widths(info, [4680, 4680])
for cell in info.rows[0].cells:
    set_cell_border(cell, top={"val": "nil"}, bottom={"val": "nil"}, start={"val": "nil"}, end={"val": "nil"})

bill, meta = info.rows[0].cells
cell_text(bill, "BILL TO", 8.5, MUTED, True)
p = bill.add_paragraph()
p.paragraph_format.space_before = Pt(3)
p.paragraph_format.space_after = Pt(1)
r = p.add_run("Vamos Taxi")
font(r, 13, CHARCOAL, True)
p = bill.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Premium transfer booking platform")
font(r, 10, MUTED)

cell_text(meta, "INVOICE DETAILS", 8.5, MUTED, True, WD_ALIGN_PARAGRAPH.RIGHT)
for label, value in (("Invoice no.", "KZ-2026-0716-01"), ("Invoice date", "16 July 2026"), ("Due date", "30 July 2026"), ("Currency", "EUR")):
    p = meta.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{label}: ")
    font(r, 9, MUTED, True)
    r = p.add_run(value)
    font(r, 10, CHARCOAL)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Line items
items = doc.add_table(rows=2, cols=3)
items.alignment = WD_TABLE_ALIGNMENT.LEFT
items.autofit = False
set_table_widths(items, [6500, 1060, 1800])
header_row = items.rows[0]
for cell, text, align in zip(header_row.cells, ["DESCRIPTION", "QTY", "AMOUNT"], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT]):
    set_cell_shading(cell, CHARCOAL)
    cell_text(cell, text, 8.5, WHITE, True, align)
    set_cell_border(cell, top={"val": "single", "sz": "4", "color": CHARCOAL}, bottom={"val": "single", "sz": "4", "color": CHARCOAL}, start={"val": "single", "sz": "4", "color": CHARCOAL}, end={"val": "single", "sz": "4", "color": CHARCOAL})

body = items.rows[1]
for cell in body.cells:
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_border(cell, top={"val": "single", "sz": "4", "color": "DDDDDD"}, bottom={"val": "single", "sz": "4", "color": "DDDDDD"}, start={"val": "single", "sz": "4", "color": "DDDDDD"}, end={"val": "single", "sz": "4", "color": "DDDDDD"})
cell_text(body.cells[0], "Vamos Taxi platform project, discovery, UX/UI direction and technical product planning", 10.5, CHARCOAL, False)
cell_text(body.cells[1], "1", 10.5, CHARCOAL, False, WD_ALIGN_PARAGRAPH.CENTER)
cell_text(body.cells[2], "€1,500.00", 10.5, CHARCOAL, True, WD_ALIGN_PARAGRAPH.RIGHT)

doc.add_paragraph().paragraph_format.space_after = Pt(4)

# Totals
totals = doc.add_table(rows=2, cols=2)
totals.alignment = WD_TABLE_ALIGNMENT.RIGHT
totals.autofit = False
set_table_widths(totals, [6660, 2700])
for row in totals.rows:
    for cell in row.cells:
        set_cell_border(cell, top={"val": "nil"}, bottom={"val": "nil"}, start={"val": "nil"}, end={"val": "nil"})

cell_text(totals.cell(0, 0), "Subtotal", 10, MUTED, False, WD_ALIGN_PARAGRAPH.RIGHT)
cell_text(totals.cell(0, 1), "€1,500.00", 10, CHARCOAL, False, WD_ALIGN_PARAGRAPH.RIGHT)
set_cell_shading(totals.cell(1, 0), YELLOW)
set_cell_shading(totals.cell(1, 1), YELLOW)
cell_text(totals.cell(1, 0), "TOTAL DUE", 11, CHARCOAL, True, WD_ALIGN_PARAGRAPH.RIGHT)
cell_text(totals.cell(1, 1), "€1,500.00", 14, CHARCOAL, True, WD_ALIGN_PARAGRAPH.RIGHT)

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# Payment note
note = doc.add_table(rows=1, cols=1)
note.alignment = WD_TABLE_ALIGNMENT.LEFT
note.autofit = False
set_table_widths(note, [9360])
cell = note.cell(0, 0)
set_cell_shading(cell, LIGHT_GRAY)
set_cell_border(cell, top={"val": "single", "sz": "4", "color": "DDDDDD"}, bottom={"val": "single", "sz": "4", "color": "DDDDDD"}, start={"val": "single", "sz": "4", "color": "DDDDDD"}, end={"val": "single", "sz": "4", "color": "DDDDDD"})
cell_text(cell, "PAYMENT TERMS", 8.5, MUTED, True)
p = cell.add_paragraph()
p.paragraph_format.space_before = Pt(3)
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Payment due by 30 July 2026. Please use the invoice number as the payment reference. Bank details will be supplied separately.")
font(r, 10, CHARCOAL)

# Footer
footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
footer_p.paragraph_format.space_before = Pt(0)
footer_p.paragraph_format.space_after = Pt(0)
r = footer_p.add_run("Thank you for your business.")
font(r, 8.5, MUTED)

doc.core_properties.title = "Vamos Taxi Invoice EUR 1,500"
doc.core_properties.author = "Koussay Zayani"
doc.core_properties.subject = "Invoice for Vamos Taxi platform project"
doc.save(OUTPUT)
print(OUTPUT)
