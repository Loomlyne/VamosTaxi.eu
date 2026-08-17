#!/usr/bin/env python3
"""Generate client-facing Scope of Work DOCX for Vamos Taxi V1."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "deliverables"
OUT.mkdir(exist_ok=True)
OUTPUT = OUT / "Vamos-Taxi-Scope-of-Work-V1.docx"

YELLOW = "FDC20B"
CHARCOAL = "1E1F1F"
LIGHT_GRAY = "F4F4F2"
MUTED = "666666"
WHITE = "FFFFFF"
BORDER = "DDDDDD"


def font(run, size=10, color=CHARCOAL, bold=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, values in kwargs.items():
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in values.items():
            element.set(qn(f"w:{key}"), str(value))


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid_cols = table._tbl.tblGrid.gridCol_lst
    for col, width in zip(grid_cols, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def set_cell_margins(cell, top=60, start=100, bottom=60, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def clear_cell(cell):
    for p in cell.paragraphs:
        p.clear()


def cell_text(cell, text, size=10, color=CHARCOAL, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    clear_cell(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    font(r, size, color, bold)
    set_cell_margins(cell)
    return p


def add_para(doc, text, size=10, color=CHARCOAL, bold=False, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(text)
    font(r, size, color, bold)
    return p


def add_heading_bar(doc, title):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_widths(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CHARCOAL)
    set_cell_margins(cell, top=80, bottom=80, start=140, end=140)
    cell_text(cell, title, 11, WHITE, True)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    # style may already have a run; clear and rewrite
    if p.runs:
        p.runs[0].text = text
        font(p.runs[0], size, CHARCOAL, False)
    else:
        r = p.add_run(text)
        font(r, size, CHARCOAL, False)
    return p


def two_col_kv(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_widths(table, [2800, 6560])
    for i, (k, v) in enumerate(rows):
        left, right = table.rows[i].cells
        fill = LIGHT_GRAY if i % 2 == 0 else WHITE
        for cell in (left, right):
            set_cell_shading(cell, fill)
            set_cell_border(
                cell,
                top={"val": "single", "sz": "4", "color": BORDER},
                bottom={"val": "single", "sz": "4", "color": BORDER},
                start={"val": "single", "sz": "4", "color": BORDER},
                end={"val": "single", "sz": "4", "color": BORDER},
            )
        cell_text(left, k, 9, MUTED, True)
        cell_text(right, v, 9.5, CHARCOAL, False)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def simple_table(doc, headers, body_rows, widths=None):
    cols = len(headers)
    widths = widths or [9360 // cols] * cols
    # fix sum
    if sum(widths) != 9360:
        widths = widths[:]
        widths[-1] += 9360 - sum(widths)
    table = doc.add_table(rows=1 + len(body_rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_widths(table, widths)
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, CHARCOAL)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "4", "color": CHARCOAL},
            bottom={"val": "single", "sz": "4", "color": CHARCOAL},
            start={"val": "single", "sz": "4", "color": CHARCOAL},
            end={"val": "single", "sz": "4", "color": CHARCOAL},
        )
        cell_text(cell, text, 8.5, WHITE, True)
    for ri, row in enumerate(body_rows):
        fill = LIGHT_GRAY if ri % 2 == 0 else WHITE
        for cell, text in zip(table.rows[ri + 1].cells, row):
            set_cell_shading(cell, fill)
            set_cell_border(
                cell,
                top={"val": "single", "sz": "4", "color": BORDER},
                bottom={"val": "single", "sz": "4", "color": BORDER},
                start={"val": "single", "sz": "4", "color": BORDER},
                end={"val": "single", "sz": "4", "color": BORDER},
            )
            cell_text(cell, text, 9, CHARCOAL, False)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.6)

    # Header
    header = doc.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.LEFT
    header.autofit = False
    set_table_widths(header, [5600, 3760])
    for cell in header.rows[0].cells:
        set_cell_border(
            cell,
            top={"val": "nil"},
            bottom={"val": "nil"},
            start={"val": "nil"},
            end={"val": "nil"},
        )
    left, right = header.rows[0].cells
    cell_text(left, "VAMOS TAXI", 16, CHARCOAL, True)
    p = left.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Platform rebuild — Scope of Work")
    font(r, 10, MUTED, False)

    cell_text(right, "DOCUMENT", 8.5, MUTED, True, WD_ALIGN_PARAGRAPH.RIGHT)
    for label, value in (
        ("Version", "1.0"),
        ("Date", "18 July 2026"),
        ("Window", "3–4 weeks / 1 month"),
        ("Status", "Client contract annex"),
    ):
        p = right.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"{label}: ")
        font(r, 9, MUTED, True)
        r = p.add_run(value)
        font(r, 9.5, CHARCOAL, False)

    # Yellow stripe
    stripe = doc.add_table(rows=1, cols=1)
    stripe.alignment = WD_TABLE_ALIGNMENT.LEFT
    stripe.autofit = False
    set_table_widths(stripe, [9360])
    cell = stripe.cell(0, 0)
    set_cell_shading(cell, YELLOW)
    set_cell_margins(cell, top=50, bottom=50, start=0, end=0)
    cell_text(cell, "", 1)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_para(
        doc,
        "This Scope of Work defines the V1 digital platform to be designed, built, integrated, and handed over for Vamos Taxi GmbH: public booking website, design system, pricing engine, payments, customer accounts, automated emails, and operations dashboard.",
        10,
        CHARCOAL,
        False,
        8,
    )

    # Parties
    add_heading_bar(doc, "1. PARTIES")
    two_col_kv(
        doc,
        [
            ("Client", "Vamos Taxi GmbH (Bleicherstrasse 16, 8953 Dietikon, Switzerland)"),
            ("Provider", "Koussay Zayani"),
            ("Product", "Scheduled airport & private transfer booking platform (Switzerland / Zurich-first)"),
            ("Model", "Book ahead for a set pickup time. Fixed price. Driver waits at place & time. Not Uber on-demand."),
            ("Reference UX", "Transfeero-level booking quality for a local operator (not a global marketplace clone)"),
        ],
    )

    # Deliverables overview
    add_heading_bar(doc, "2. DELIVERABLES OVERVIEW")
    simple_table(
        doc,
        ["Area", "What you get"],
        [
            ("Design", "Mobile-first premium UI, design system in code, booking-first layouts, admin UI"),
            ("Public website", "Homepage, multi-step booking, map pin/search, quote, vehicles, extras, coupons, checkout"),
            ("Accounts", "Customer sign-in, profile, booking history, manage booking (+ guest checkout path)"),
            ("Pricing engine", "Fixed routes + distance-based calculation, surcharges, extras, coupon discounts, price snapshot"),
            ("Payments", "Stripe standard checkout + webhooks (not Stripe Connect)"),
            ("Emails", "Resend: confirmation, voucher, status updates, pre-trip reminders"),
            ("Maps", "Mapbox: search, click-to-set pickup/destination, planned route for pricing (no live GPS)"),
            ("Admin dashboard", "Bookings, calendar, assign driver/vehicle, manual bookings, pricing & coupons admin, CSV export"),
            ("Launch", "Production deploy, domain cutover support, source code ownership, admin handover"),
        ],
        [2600, 6760],
    )

    # Public site detail
    add_heading_bar(doc, "3. PUBLIC WEBSITE (CUSTOMER)")
    for t in [
        "Homepage with booking widget above the fold (mobile-first, excellent on desktop).",
        "Pickup & destination via address search and map click (Mapbox).",
        "Date & time selection for scheduled transfers (e.g. book today for tomorrow).",
        "Passengers, luggage, one-way or return journey.",
        "Server-side quote: distance, duration, eligible vehicle classes with fixed prices.",
        "Passenger details, flight number field, configurable extras.",
        "Coupon / voucher codes at checkout.",
        "Stripe payment (methods available on Client Stripe account; TWINT if eligible).",
        "Confirmation page and voucher content.",
        "Customer accounts: registration/login, booking history, open booking.",
        "Guest checkout with email; option to claim booking into an account.",
        "Manage booking within Client cancellation / modification policy.",
        "About, Contact, FAQ + legal page structure (Terms, Privacy, Imprint, Cancellation).",
        "English + German interface for V1; CHF primary currency (unless kickoff confirms otherwise).",
        "Initial SEO structure for key airport / fixed routes once Client provides the list.",
    ]:
        add_bullet(doc, t)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Pricing
    add_heading_bar(doc, "4. PRICING & CALCULATION ENGINE")
    add_para(doc, "All prices are calculated on the server. The browser cannot invent fares.", 10, MUTED, False, 4)
    simple_table(
        doc,
        ["Rule type", "Behaviour"],
        [
            ("Fixed routes", "Configured A→B price per vehicle class overrides calculated price"),
            ("Calculated", "Base + distance rate + surcharges + vehicle factor + extras − discount"),
            ("Surcharges", "Airport, night/time, waiting — as supplied by Client"),
            ("Coupons", "% or fixed CHF, min order, expiry, max uses"),
            ("Snapshot", "Full breakdown stored on each booking for stable history"),
            ("Guards", "Min advance booking, capacity vs passengers/luggage"),
        ],
        [2400, 6960],
    )

    # Integrations
    add_heading_bar(doc, "5. SYSTEM INTEGRATIONS")
    simple_table(
        doc,
        ["System", "Role in V1"],
        [
            ("Vercel", "Hosting and production deployment"),
            ("Supabase", "Database, authentication, security (RLS), storage"),
            ("Mapbox", "Maps, geocoding, directions for distance/time"),
            ("Stripe", "Online payments + payment webhooks (standard merchant account)"),
            ("Resend", "Transactional emails (confirmations, reminders, vouchers)"),
        ],
        [2200, 7160],
    )
    add_para(
        doc,
        "Client owns all third-party accounts and pays ongoing usage fees (hosting, database, maps, email, Stripe fees). Provider configures integrations during the build and hands access over.",
        9.5,
        MUTED,
        False,
        8,
    )

    # Automation
    add_heading_bar(doc, "6. AUTOMATION (SCHEDULED MODEL)")
    for t in [
        "Payment success → booking confirmed → voucher email to customer.",
        "New booking visible in admin; optional ops notification email.",
        "Pre-trip reminder email before pickup time.",
        "Customer emails on key status changes (confirmed, cancelled, etc.).",
        "Cancellation window enforced per Client policy.",
        "No live driver GPS or “car arriving in 3 minutes” — driver waits at the booked time.",
    ]:
        add_bullet(doc, t)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Admin
    add_heading_bar(doc, "7. OPERATIONS DASHBOARD")
    for t in [
        "Secure admin login with role checks.",
        "Bookings list (search/filter) sorted by pickup date/time.",
        "Day / calendar style view of transfers.",
        "Booking detail: passenger, route, price breakdown, payment, notes.",
        "Manual booking creation for phone customers.",
        "Status workflow (e.g. new → confirmed → assigned → completed / cancelled).",
        "Manual assignment of driver and vehicle (records only — no driver app).",
        "Driver and vehicle directory.",
        "Admin for fixed routes, pricing rules, extras, and coupons.",
        "Customer records linked to bookings.",
        "CSV export and audit trail for important changes.",
    ]:
        add_bullet(doc, t)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Out of scope
    add_heading_bar(doc, "8. OUT OF SCOPE (V1)")
    add_para(doc, "Not included unless agreed in a written change order:", 10, MUTED, False, 4)
    for t in [
        "Native iOS/Android apps (customer or driver).",
        "Driver mobile app or driver self-service portal.",
        "Live continuous GPS tracking of drivers/vehicles.",
        "On-demand “ride now” matching or surge pricing.",
        "Global multi-country marketplace or multi-fleet SaaS.",
        "Automatic nearest-driver dispatch algorithms.",
        "Stripe Connect multi-party payouts / driver wallets / payroll.",
        "Hotel or affiliate partner portals.",
        "Full historical migration from Freshpage (export help only if Client can export).",
        "Legal drafting of Swiss T&Cs/privacy (Client or lawyer supplies final text).",
        "Photo shoots, paid ads, large-scale SEO copywriting, social media management.",
        "Ongoing 24/7 support retainer (can be contracted separately after launch).",
    ]:
        add_bullet(doc, t)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Timeline
    add_heading_bar(doc, "9. TIMELINE (3–4 WEEKS)")
    add_para(
        doc,
        "Kickoff = Client accepts this Scope and supplies Blocking Inputs (Section 10). Target: finish within one calendar month.",
        9.5,
        MUTED,
        False,
        4,
    )
    simple_table(
        doc,
        ["Week", "Focus", "Outcome"],
        [
            ("Week 1", "Foundation + design + quote", "Staging app; booking widget; Mapbox; route & price; vehicle classes"),
            ("Week 2", "Checkout + accounts + email", "Pay path; coupons; auth + history; confirmation emails"),
            ("Week 3", "Operations dashboard", "Run daily ops: bookings, assign, statuses, pricing admin"),
            ("Week 4", "Polish + launch", "Legal/content, DE+EN, QA, production domain, acceptance"),
        ],
        [1400, 2600, 5360],
    )
    add_para(
        doc,
        "Delays in Client inputs, Stripe readiness, DNS access, or legal copy move the end date day-for-day.",
        9.5,
        MUTED,
        False,
        8,
    )

    # Client inputs
    add_heading_bar(doc, "10. CLIENT RESPONSIBILITIES")
    add_para(doc, "Blocking inputs (for pricing & rules):", 10, CHARCOAL, True, 3)
    for t in [
        "Company legal details for imprint.",
        "Service area (cities / airports).",
        "Vehicle classes: name, max passengers, max luggage.",
        "Pricing: base, per-km or class rates, minimum fare, fixed routes, airport/night/extras.",
        "Policies: min advance booking, free cancel window, refunds, waiting rules.",
        "Confirm CHF + English/German (recommended defaults).",
    ]:
        add_bullet(doc, t)
    add_para(doc, "Accounts & access:", 10, CHARCOAL, True, 6, 4)
    for t in [
        "Domain DNS access for production.",
        "Stripe business account.",
        "Email domain DNS for Resend.",
        "Ownership of Vercel, Supabase, Mapbox projects at handover.",
        "Feedback on staging within 48 hours so the schedule holds.",
        "Final legal texts (Terms, Privacy, Cancellation).",
        "Logo, brand fonts if required, vehicle photos.",
    ]:
        add_bullet(doc, t)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Acceptance
    add_heading_bar(doc, "11. ACCEPTANCE CRITERIA")
    for t in [
        "Customer completes scheduled booking on mobile and desktop (quote → pay/offline → confirmation).",
        "Prices match configured rules for a fixed-route test and a calculated-route test.",
        "Valid coupon reduces total correctly.",
        "Stripe payment (test or live) marks booking paid via webhook.",
        "Customer receives confirmation/voucher email.",
        "Account shows booking history.",
        "Admin can view booking, assign driver/vehicle, change status, create manual booking.",
        "Admin can edit core pricing rules and coupons.",
        "Agreed content/legal pages published.",
        "Production live on agreed domain; credentials and code handed over.",
    ]:
        add_bullet(doc, t)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Commercial / legal
    add_heading_bar(doc, "12. COMMERCIAL, IP & SUPPORT")
    two_col_kv(
        doc,
        [
            ("Delivery fee", "As agreed on the commercial quotation / invoice attached to this Scope"),
            ("Third-party costs", "Paid by Client (hosting, DB, maps, email, Stripe fees, domain)"),
            ("IP", "Upon full payment of agreed fees, Client owns custom V1 source code delivered for Vamos Taxi"),
            ("Third-party licences", "Stripe, Mapbox, Supabase, Vercel, Resend, fonts remain under their own terms"),
            ("Changes", "Work outside this Scope needs written change order (scope, time, fee impact)"),
            ("Bug fix window", "14 days after acceptance for defects in accepted V1 flows (not new features)"),
            ("Validity", "Scope valid 30 days from document date unless accepted earlier"),
        ],
    )

    # Signature
    add_heading_bar(doc, "13. ACCEPTANCE SIGNATURES")
    add_para(
        doc,
        "By signing, Client confirms this Scope describes the V1 project to be delivered in the stated window, subject to timely Client inputs and the exclusions listed.",
        9.5,
        MUTED,
        False,
        8,
    )

    sig = doc.add_table(rows=5, cols=2)
    sig.alignment = WD_TABLE_ALIGNMENT.LEFT
    sig.autofit = False
    set_table_widths(sig, [4680, 4680])
    labels = [
        ("CLIENT", "PROVIDER"),
        ("Name: Ben Othman Houssein / authorised signatory", "Name: Koussay Zayani"),
        ("Company: Vamos Taxi GmbH", "Company: —"),
        ("Signature: ___________________________", "Signature: ___________________________"),
        ("Date: _______________________________", "Date: _______________________________"),
    ]
    for i, (a, b) in enumerate(labels):
        ca, cb = sig.rows[i].cells
        for cell in (ca, cb):
            set_cell_border(
                cell,
                top={"val": "nil"},
                bottom={"val": "nil"},
                start={"val": "nil"},
                end={"val": "nil"},
            )
            set_cell_margins(cell, top=40, bottom=40, start=40, end=40)
        bold = i == 0
        color = MUTED if i == 0 else CHARCOAL
        size = 9 if i == 0 else 9.5
        cell_text(ca, a, size, color, bold)
        cell_text(cb, b, size, color, bold)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.LEFT
    box.autofit = False
    set_table_widths(box, [9360])
    cell = box.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_border(
        cell,
        top={"val": "single", "sz": "4", "color": BORDER},
        bottom={"val": "single", "sz": "4", "color": BORDER},
        start={"val": "single", "sz": "4", "color": BORDER},
        end={"val": "single", "sz": "4", "color": BORDER},
    )
    cell_text(cell, "ONE-LINE SUMMARY FOR THE CLIENT", 8.5, MUTED, True)
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(
        "In 3–4 weeks you receive a full scheduled-transfer platform: premium website + design, Mapbox quote, pricing engine, Stripe pay, coupons, customer accounts & history, automated emails, and a dispatch dashboard — not apps, not live GPS, not Uber."
    )
    font(r, 10, CHARCOAL, False)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = fp.add_run("Vamos Taxi — Scope of Work V1 · Confidential · Provider: Koussay Zayani")
    font(r, 8, MUTED, False)

    doc.core_properties.title = "Vamos Taxi Scope of Work V1"
    doc.core_properties.author = "Koussay Zayani"
    doc.core_properties.subject = "Statement of Work for Vamos Taxi platform rebuild"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
